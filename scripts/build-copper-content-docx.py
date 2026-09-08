from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DATA_SQL = ROOT / "springboot" / "src" / "main" / "resources" / "data.sql"
OUTPUT = ROOT / "output" / "3d-copper-content" / "小铜人内容中心_穴位与故事内容汇编.docx"
MASCOT = ROOT / "vue3" / "src" / "assets" / "copper-detective-2026" / "detective-guide.png"

# compact_reference_guide preset + named "Xinglin brand" color/font override.
DEEP_TEAL = "1F5A50"
MID_TEAL = "2F7D68"
GOLD = "B47A2E"
CREAM = "FFF8E8"
PALE_GREEN = "E8F2EE"
PALE_BLUE = "E8EEF5"
INK = "263B35"
MUTED = "6C7A75"
WHITE = "FFFFFF"
BRAND_FONT = "Noto Sans SC"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, size=10.5, bold=False, color=INK, italic=False) -> None:
    run.font.name = BRAND_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BRAND_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), BRAND_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), BRAND_FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)


def add_text(paragraph, text, size=10.5, bold=False, color=INK, italic=False):
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return run


def style_paragraph(paragraph, before=0, after=6, line=1.25, keep=False) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    add_text(p, text, 10.5)
    style_paragraph(p, after=4, line=1.25)
    return p


def add_callout(doc, label, text, fill=CREAM):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    add_text(p, f"{label}  ", 10.5, True, GOLD)
    add_text(p, text, 10.5, False, INK)
    style_paragraph(p, after=0, line=1.25)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_sql_values(segment: str) -> list[str | None]:
    values: list[str | None] = []
    token: list[str] = []
    quoted = False
    i = 0
    while i < len(segment):
        char = segment[i]
        if quoted:
            if char == "'":
                if i + 1 < len(segment) and segment[i + 1] == "'":
                    token.append("'")
                    i += 2
                    continue
                quoted = False
            elif char == "\\" and i + 1 < len(segment):
                token.append(segment[i + 1])
                i += 2
                continue
            else:
                token.append(char)
        else:
            if char == "'":
                quoted = True
            elif char == ",":
                raw = "".join(token).strip()
                values.append(None if raw.upper() == "NULL" else raw)
                token = []
            else:
                token.append(char)
        i += 1
    raw = "".join(token).strip()
    values.append(None if raw.upper() == "NULL" else raw)
    return values


def load_acupoints() -> list[dict]:
    rows = []
    for line in DATA_SQL.read_text(encoding="utf-8").splitlines():
        if not line.startswith("INSERT INTO `acupoint_knowledge`"):
            continue
        segment = line.split(" VALUES (", 1)[1].rsplit(") ON DUPLICATE", 1)[0]
        values = parse_sql_values(segment)
        if len(values) != 20:
            raise RuntimeError(f"Unexpected acupoint column count: {len(values)}")
        rows.append(dict(zip([
            "code", "point_number", "name", "pinyin", "meridian_code", "meridian_name",
            "meridian_english", "model_id", "body_area", "standard_location", "child_location",
            "child_description", "safety_tip", "source_name", "source_link", "position_x",
            "position_y", "position_z", "sort_order", "enabled"
        ], values)))
    if len(rows) != 361:
        raise RuntimeError(f"Expected 361 acupoints, got {len(rows)}")
    return rows


STORIES = [
    {
        "title": "失踪竹简案", "code": "missing-bamboo", "summary": "风把竹简藏到了书架后，请跟着竹叶和铜片找到它。",
        "pages": [
            ("竹简架空了一格", "清晨，小铜人老师发现一卷竹简不见了。桌边有一枚铜片，窗边有一片竹叶。", "先看现场，再猜答案。"),
            ("小风留下了线索", "竹叶轻轻摇，铜片滚到桌角。原来，窗口刚才吹进一阵风。", "把竹叶、铜片和空架连起来。"),
            ("竹简在书架后", "小侦探在书架后找到了竹简。它是被风吹落的，没有人把它拿走。", "好推理要有线索帮忙。"),
        ],
        "clues": [("空竹简架", "这里原来放着竹简。"), ("窗边竹叶", "竹叶告诉我们：刚才有风。"), ("圆圆铜片", "铜片被风推到了桌角。")],
        "reasoning": ("竹简去了哪里？", "被风吹到书架后 / 自己跑出了门", "被风吹到书架后", "竹叶和铜片都指向书架后。"),
        "safety": ("看到穴位知识时怎么做？", "只观察，请老师讲解", "我们只学文化知识，不做身体操作。"),
        "reward": "竹简碎片 ×3", "acupoints": "无"
    },
    {
        "title": "会考试的小铜人", "code": "exam-copper-man", "summary": "跟着王惟一走进北宋课堂，看看铜人怎样帮大家学习。",
        "pages": [
            ("王惟一的铜人课堂", "北宋时，王惟一与工匠制作了铜人模型。模型上有许多小标记。", "铜人是学习模型，不是玩具。"),
            ("星点变成学习题", "学生先记名字，再在铜人上找标记。找对了，就完成一道古代学习题。", "我们也用3D铜人只做观察。"),
            ("老铜人来到新课堂", "今天，铜人变成了屏幕里的3D文化地图。我们可以转一转、看一看。", "文化知识要和安全规则一起学。"),
        ],
        "clues": [("铜人标记", "小标记帮学生认识名字和路线。"), ("课堂记录", "记录里写着北宋的铜人故事。")],
        "reasoning": ("古人为什么制作铜人？", "帮助观察和学习 / 让铜人自己看病", "帮助观察和学习", "铜人是认识身体地图的教学模型。"),
        "safety": ("我们怎样使用3D铜人？", "转动模型观察", "只观察、只学习。"),
        "reward": "铜片 ×2", "acupoints": "LU-1、LU-5、LU-9"
    },
    {
        "title": "星河断线案", "code": "broken-star-river", "summary": "三颗文化星迷了路，请把它们送回同一条经络星河。",
        "pages": [
            ("星河忽然暗了", "中府星、尺泽星和太渊星都亮着，它们之间的光线却不见了。", "看看三颗星的经络名称。"),
            ("三颗星找到同一条路", "它们都属于手太阴肺经。名字像路牌，帮它们排到同一条路线上。", "经络可以当成文化地图的路线名。"),
            ("星河又亮了", "小侦探按照路线名称连好星点。经络星河又变得明亮了。", "今天记住的是名字和路线。"),
        ],
        "clues": [("相同路线名", "三颗星的档案都写着手太阴肺经。"), ("星点编号", "编号帮我们整理文化地图。")],
        "reasoning": ("哪个线索能把星点连起来？", "相同的经络名称 / 卡片的颜色", "相同的经络名称", "经络名称是这张文化地图的路线牌。"),
        "safety": ("学习路线时要记住什么？", "只看图和模型", "我们只认识文化路线。"),
        "reward": "经络星砂 ×3", "acupoints": "LU-1、LU-5、LU-9"
    },
    {
        "title": "安全铃铛失声案", "code": "silent-safety-bell", "summary": "安全铃铛不响了，请找回“只观察”的金色声音。",
        "pages": [
            ("铃铛怎么不响了", "小药师听到一句“跟着视频试一试”。安全铃铛立刻安静了。", "这句话安全吗？"),
            ("三张安全卡", "第一张写“只观察”，第二张写“问成人”，第三张写“不模仿”。", "三张卡都是安全线索。"),
            ("金色声音回来了", "小侦探大声说：“只观察、不模仿，有问题找老师和家长！”铃铛又响了。", "记住这句安全口令。"),
        ],
        "clues": [("只观察卡", "看3D模型和文化图卡是安全的。"), ("问成人卡", "不舒服或有疑问时，找家长、老师或医生。"), ("不模仿卡", "不照着图片或视频做身体操作。")],
        "reasoning": ("哪句话能让安全铃铛响起来？", "只观察，有问题问成人 / 跟着视频自己试", "只观察，有问题问成人", "学文化知识时，安全规则永远排在第一位。"),
        "safety": ("身体不舒服时应该怎么做？", "告诉家长并询问医生", "不舒服要告诉可信任的成人。"),
        "reward": "安全铃铛 ×1", "acupoints": "无"
    },
]


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BRAND_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BRAND_FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    specs = {
        "Heading 1": (16, DEEP_TEAL, 18, 10),
        "Heading 2": (13, MID_TEAL, 14, 7),
        "Heading 3": (12, GOLD, 10, 5),
    }
    for style_name, (size, color, before, after) in specs.items():
        style = styles[style_name]
        style.font.name = BRAND_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BRAND_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Bullet 2"):
        style = styles[style_name]
        style.font.name = BRAND_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BRAND_FONT)
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(paragraph, "小铜人内容中心 · ", 8.5, False, MUTED)
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run_font(run, 8.5, color=MUTED)


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        paragraph._p.remove(child)


def configure_section(section, landscape=False) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
    else:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    clear_paragraph(header)
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(header, "3D 小铜人馆内容资产汇编", 8.5, True, MID_TEAL)
    footer = section.footer.paragraphs[0]
    clear_paragraph(footer)
    add_page_field(footer)


def add_cover(doc: Document, acupoints: list[dict]) -> None:
    p = doc.add_paragraph()
    style_paragraph(p, after=18)
    add_text(p, "人体探案实验室 · 内容资产", 10.5, True, GOLD)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_paragraph(p, after=10, keep=True)
    add_text(p, "小铜人内容中心", 29, True, DEEP_TEAL)
    p = doc.add_paragraph()
    style_paragraph(p, after=22)
    add_text(p, "穴位儿童文案、探案故事与后台运营说明汇编", 15, False, MID_TEAL)

    if MASCOT.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run()
        run.add_picture(str(MASCOT), width=Inches(2.15))
        style_paragraph(p, after=4)

    modeled = sum(1 for point in acupoints if point["position_x"] is not None and point["position_y"] is not None and point["position_z"] is not None)
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [3120, 3120, 3120])
    metrics = [("361", "穴位文化星"), (str(modeled), "3D 已就绪"), (str(len(STORIES)), "完整探案故事")]
    for cell, (value, label) in zip(table.rows[0].cells, metrics):
        set_cell_shading(cell, PALE_GREEN)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, value, 22, True, DEEP_TEAL)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, label, 9, True, MUTED)

    doc.add_paragraph()
    add_callout(doc, "永久安全边界", "本汇编仅用于儿童中医药文化教育。儿童端不提供诊断、治疗功效、自我取穴、按揉、刺激或针刺指导。", CREAM)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_paragraph(p, before=18, after=0)
    add_text(p, "版本：2026-08-26 · 适龄：5–8 岁 · 数据源：acupoint_knowledge", 9.5, False, MUTED)
def add_overview(doc: Document) -> None:
    heading = add_heading(doc, "一、内容中心总览", 1)
    heading.paragraph_format.page_break_before = True
    add_callout(doc, "建设结论", "穴位由 acupoint_knowledge 统一供给；故事由管理员编排、审核、发布和推送。旧 xuewei 数据保留兼容，但不再作为儿童端内容源。", PALE_GREEN)
    add_heading(doc, "1. 内容分层", 2)
    add_bullet(doc, "专业层：编号、名称、拼音、经络、标准定位和专业来源，仅用于管理员核对与内容溯源。")
    add_bullet(doc, "儿童层：星点称呼、模型观察提示、身体区域、文化小档案、记忆线索和安全提醒。")
    add_bullet(doc, "3D 规则：只有 X/Y/Z 坐标完整且通过校验的穴位，才可进入每日探案任务池。")
    add_heading(doc, "2. 儿童内容写作规则", 2)
    for item in [
        "短句优先；一个页面只承担一件观察或推理任务。",
        "用“文化星、路线牌、地图标签、档案卡”等比喻解释穴位与经络。",
        "不出现治疗承诺，不引导儿童在自己或他人身体上寻找、按揉、刺激或针刺。",
        "每条穴位都保留来源，并固定显示“只观察、不操作、有问题问成人”的安全提醒。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3. 审核发布工作流", 2)
    workflow = doc.add_table(rows=1, cols=5)
    set_table_geometry(workflow, [1872] * 5)
    for cell, text in zip(workflow.rows[0].cells, ["草稿", "待审核", "已审核", "立即/定时发布", "已归档"]):
        set_cell_shading(cell, PALE_BLUE if text not in ("已审核", "立即/定时发布") else PALE_GREEN)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, text, 9.5, True, DEEP_TEAL)
    p = doc.add_paragraph()
    add_text(p, "每次提交、审核、发布和归档均记录操作人、时间与版本；重复发布不会重复生成版本、推送或奖励。", 10.5)
    style_paragraph(p, before=5, after=7)

    add_heading(doc, "4. 管理端与儿童端功能", 2)
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [4680, 4680])
    for cell, title in zip(table.rows[0].cells, ["管理员：小铜人内容中心", "儿童端：人体探案实验室"]):
        set_cell_shading(cell, DEEP_TEAL)
        p = cell.paragraphs[0]
        add_text(p, title, 10.5, True, WHITE)
    content = [
        ["穴位星图：完整度、3D 就绪、儿童文案、来源、状态\n故事工坊：场景、旁白、图片、线索、推理、安全题、奖励\n发布队列：待审核、定时、失败、撤回\n推送记录：目标内容与已读统计", "故事馆读取后台已发布内容\n首页/导航显示新故事与新星点角标\n消息点击后直达故事或穴位图鉴\n故事进度与结案状态保存到服务端"],
    ]
    for left, right in content:
        row = table.add_row()
        for cell, text in zip(row.cells, [left, right]):
            for index, line in enumerate(text.split("\n")):
                p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
                add_text(p, line, 9)
                style_paragraph(p, after=1, line=1.05)
def add_story_section(doc: Document) -> None:
    heading = add_heading(doc, "二、首批儿童探案故事", 1)
    heading.paragraph_format.page_break_before = True
    p = doc.add_paragraph()
    add_text(p, "每个故事固定完成“发生事件—观察场景—收集线索—做出判断—揭晓答案—安全回顾”的完整闭环。", 10.5)
    style_paragraph(p, after=10)

    for index, story in enumerate(STORIES, 1):
        if index > 1:
            spacer = doc.add_paragraph()
            spacer.paragraph_format.page_break_before = True
            spacer.paragraph_format.space_after = Pt(6)
        story_heading = add_heading(doc, f"{index}. 《{story['title']}》", 2)
        meta = doc.add_table(rows=3, cols=2)
        set_table_geometry(meta, [1875, 7485])
        pairs = [("故事编号", story["code"]), ("推送简介", story["summary"]), ("关联穴位 / 奖励", f"{story['acupoints']} / {story['reward']}")]
        for row, (label, value) in zip(meta.rows, pairs):
            set_cell_shading(row.cells[0], PALE_BLUE)
            add_text(row.cells[0].paragraphs[0], label, 9.5, True, DEEP_TEAL)
            add_text(row.cells[1].paragraphs[0], value, 9.5)

        add_heading(doc, "故事场景", 3)
        scenes = doc.add_table(rows=1, cols=4)
        set_table_geometry(scenes, [750, 1900, 5000, 1710])
        for cell, header in zip(scenes.rows[0].cells, ["序号", "场景", "儿童旁白", "观察提示"]):
            set_cell_shading(cell, DEEP_TEAL)
            add_text(cell.paragraphs[0], header, 9, True, WHITE)
        repeat_table_header(scenes.rows[0])
        for scene_index, (title, text, tip) in enumerate(story["pages"], 1):
            row = scenes.add_row()
            for cell, value in zip(row.cells, [scene_index, title, text, tip]):
                add_text(cell.paragraphs[0], value, 9)
                style_paragraph(cell.paragraphs[0], after=0, line=1.15)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            keep_row_together(row)

        add_heading(doc, "线索、推理与安全回顾", 3)
        clues = "；".join(f"{name}：{description}" for name, description in story["clues"])
        reasoning = story["reasoning"]
        safety = story["safety"]
        detail = doc.add_table(rows=3, cols=2)
        set_table_geometry(detail, [1875, 7485])
        rows = [
            ("线索", clues),
            ("推理题", f"{reasoning[0]}  选项：{reasoning[1]}  正确答案：{reasoning[2]}。{reasoning[3]}"),
            ("安全题", f"{safety[0]}  正确答案：{safety[1]}。安全口令：{safety[2]}"),
        ]
        for row, (label, value) in zip(detail.rows, rows):
            set_cell_shading(row.cells[0], CREAM if label == "安全题" else PALE_BLUE)
            add_text(row.cells[0].paragraphs[0], label, 9.5, True, GOLD if label == "安全题" else DEEP_TEAL)
            add_text(row.cells[1].paragraphs[0], value, 9.3)
            style_paragraph(row.cells[1].paragraphs[0], after=0, line=1.18)
def add_acupoint_appendix(doc: Document, acupoints: list[dict]) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section, landscape=True)
    add_heading(doc, "三、361 条穴位内容清单", 1)
    add_callout(doc, "阅读说明", "“3D可探案”表示坐标完整，可进入每日探案；“仅图鉴”表示当前无完整坐标，只在百穴图鉴展示。标准定位仅供管理员核对，不向儿童展示。", CREAM)

    grouped: dict[tuple[str, str], list[dict]] = {}
    for point in acupoints:
        grouped.setdefault((point["meridian_code"], point["meridian_name"]), []).append(point)

    widths = [1050, 1600, 1250, 2600, 3000, 3000, 1468]
    for group_index, ((meridian_code, meridian_name), points) in enumerate(grouped.items(), 1):
        heading = add_heading(doc, f"{meridian_name}（{meridian_code}）· {len(points)} 穴", 2)
        heading.paragraph_format.page_break_before = group_index > 1
        table = doc.add_table(rows=1, cols=7)
        set_table_geometry(table, widths)
        headers = ["编号", "名称/拼音", "身体区域", "专业定位/来源", "儿童观察提示", "文化记忆线索", "状态"]
        for cell, header in zip(table.rows[0].cells, headers):
            set_cell_shading(cell, DEEP_TEAL)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_text(p, header, 8.2, True, WHITE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        repeat_table_header(table.rows[0])
        for point in points:
            modeled = all(point[key] is not None for key in ("position_x", "position_y", "position_z"))
            professional = point["standard_location"] or "项目现有穴位目录暂未提供标准定位。"
            professional += f"\n来源：{point['source_name'] or '待补充'}"
            status = "3D可探案" if modeled else "仅图鉴"
            values = [
                point["code"],
                f"{point['name']}\n{point['pinyin'] or '拼音待补充'}",
                point["body_area"],
                professional,
                point["child_location"],
                point["child_description"],
                status,
            ]
            row = table.add_row()
            if not modeled:
                for cell in row.cells:
                    set_cell_shading(cell, "FBF8EF")
            for cell, value in zip(row.cells, values):
                p = cell.paragraphs[0]
                for part_index, part in enumerate(str(value).split("\n")):
                    target = p if part_index == 0 else cell.add_paragraph()
                    add_text(target, part, 7.4 if part_index == 0 else 7.0, bold=(cell is row.cells[0] or (cell is row.cells[6] and part_index == 0)), color=DEEP_TEAL if cell in (row.cells[0], row.cells[6]) else INK)
                    style_paragraph(target, after=1, line=1.05)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            keep_row_together(row)
        p = doc.add_paragraph()
        add_text(p, f"本经络安全提醒统一为：{points[0]['safety_tip']}", 8, False, GOLD)
        style_paragraph(p, before=3, after=0, line=1.1)


def add_source_note(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section, landscape=False)
    add_heading(doc, "四、来源与使用说明", 1)
    add_bullet(doc, "专业定位与来源字段沿用项目现有穴位目录及 GB/T 12346-2021《经穴名称与定位》数据。")
    add_bullet(doc, "儿童文案为 5–8 岁文化教育表述，不构成医疗建议、诊断或治疗说明。")
    add_bullet(doc, "无完整 3D 坐标的穴位只能发布到图鉴，不进入每日探案任务池。")
    add_bullet(doc, "故事与穴位的发布状态、版本、提交人、审核人和发布时间由后台数据库留痕。")
    add_callout(doc, "安全口令", "只观察、不模仿；不舒服或有疑问时，告诉家长、老师或医生。", CREAM)


def main() -> None:
    acupoints = load_acupoints()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = False
    configure_styles(doc)
    configure_section(doc.sections[0], landscape=False)
    add_cover(doc, acupoints)
    add_overview(doc)
    add_story_section(doc)
    add_acupoint_appendix(doc, acupoints)
    add_source_note(doc)

    props = doc.core_properties
    props.title = "小铜人内容中心：穴位与故事内容汇编"
    props.subject = "3D 小铜人馆 361 条穴位儿童内容、四个探案故事与管理说明"
    props.author = "小铜人内容中心"
    props.keywords = "小铜人,穴位文化,儿童故事,内容管理,安全教育"
    props.comments = "面向 5–8 岁儿童的文化教育内容，不提供诊疗或身体操作指导。"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
